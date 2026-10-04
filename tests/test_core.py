import os

import pytest
from openpyxl import load_workbook

from renal_biopsy.db import ConflictError, Database
from renal_biopsy.export import export_records, pick_lab_row
from renal_biopsy.utils import calc_age, calc_bmi, normalize_date, to_number


@pytest.fixture
def db(tmp_path):
    d = Database(str(tmp_path / "test.db"))
    yield d
    d.close()


def make_record(db, chart, date, gender="Male", birth="1970-05-01", dx=("IgA nephropathy",)):
    rec = db.empty_record(chart)
    rec["patient"].update(birth_date=birth, gender=gender)
    rec["biopsy"]["data"]["biopsy_date"] = date
    rec["biopsy"]["diagnoses"] = list(dx)
    rec["labs"] = [
        {"panel": "chem", "lab_date": "2020-01-01", "values": {"Cr": "1.0"}},
        {"panel": "chem", "lab_date": date, "values": {"Cr": "2.5", "BUN": "30"}},
    ]
    return rec


def test_utils():
    assert normalize_date("2024/1/5") == "2024-01-05"
    assert normalize_date("20240105") == "2024-01-05"
    assert normalize_date("113/01/05") == "2024-01-05"
    assert normalize_date("") == ""
    assert normalize_date("2024-02-30") is None
    assert calc_age("1970-05-01", "2024-04-30") == "53"
    assert calc_age("1970-05-01", "2024-05-01") == "54"
    assert calc_bmi("70", "170") == "24.2"
    assert to_number("3") == 3 and to_number("2.50") == 2.5 and to_number("<0.5") == "<0.5"


def test_create_and_reload(db):
    rec = make_record(db, "A001", "2024-03-01")
    rec["images"] = [{"filename": "us.png", "data": b"\x89PNG"}]
    bid = db.save_record(rec, "tester")
    assert db.get_patient("A001")["gender"] == "Male"
    loaded = db.load_record(bid)
    assert loaded["biopsy"]["data"]["biopsy_date"] == "2024-03-01"
    assert loaded["biopsy"]["data"]["pe_consciousness"] == "clear"  # 預設值
    assert len(loaded["labs"]) == 2
    assert loaded["images"][0]["data"] == b"\x89PNG"

    # 修改並刪除圖片
    loaded["biopsy"]["data"]["path_lm"] = "mesangial proliferation"
    loaded["images"] = []
    db.save_record(loaded, "tester")
    again = db.load_record(bid)
    assert again["biopsy"]["data"]["path_lm"] == "mesangial proliferation"
    assert again["images"] == []
    assert again["biopsy"]["version"] == 2


def test_multiple_biopsies_per_patient(db):
    db.save_record(make_record(db, "A001", "2020-01-01"), "u")
    db.save_record(make_record(db, "A001", "2024-01-01"), "u")
    assert [b["biopsy_date"] for b in db.list_biopsies("A001")] == ["2020-01-01", "2024-01-01"]


def test_conflict_detection(db):
    bid = db.save_record(make_record(db, "A001", "2024-03-01"), "u")
    a = db.load_record(bid)
    b = db.load_record(bid)
    a["biopsy"]["data"]["path_lm"] = "A"
    db.save_record(a, "userA")
    b["biopsy"]["data"]["path_lm"] = "B"
    with pytest.raises(ConflictError):
        db.save_record(b, "userB")
    assert db.load_record(bid)["biopsy"]["data"]["path_lm"] == "A"


def test_patient_conflict_only_when_patient_changed(db):
    bid = db.save_record(make_record(db, "A001", "2024-03-01"), "u")
    a = db.load_record(bid)
    other = db.empty_record("A001")
    other["biopsy"]["data"]["biopsy_date"] = "2025-01-01"
    other["patient"]["past_history"] = "DM"
    db.save_record(other, "userB")
    # a 沒改病人資料 → 可以存
    a["biopsy"]["data"]["path_if"] = "IgA 3+"
    db.save_record(a, "userA")
    # a 改了病人資料，但版本已過期 → 衝突
    a["patient"]["gender"] = "Female"
    with pytest.raises(ConflictError):
        db.save_record(a, "userA")


def test_new_patient_race(db):
    r1 = make_record(db, "A001", "2024-03-01")
    r2 = make_record(db, "A001", "2024-03-02", gender="Female")
    db.save_record(r1, "u1")
    with pytest.raises(ConflictError):
        db.save_record(r2, "u2")


def test_locks(db, tmp_path):
    bid = db.save_record(make_record(db, "A001", "2024-03-01"), "u")
    assert db.acquire_lock(bid, "pc1") == (True, "pc1")
    other = Database(db.path)
    assert other.acquire_lock(bid, "pc2") == (False, "pc1")
    db.release_lock(bid, "pc1")
    assert other.acquire_lock(bid, "pc2") == (True, "pc2")
    # 過期的鎖可被接手
    other.conn.execute("UPDATE locks SET heartbeat='2000-01-01 00:00:00'")
    assert db.acquire_lock(bid, "pc1") == (True, "pc1")
    other.close()


def test_delete(db):
    bid = db.save_record(make_record(db, "A001", "2024-03-01"), "u")
    db.delete_biopsy(bid, "admin")
    assert db.get_patient("A001") is None


def test_query_and_export(db, tmp_path):
    db.save_record(make_record(db, "A001", "2023-06-01", "Male"), "u")
    db.save_record(make_record(db, "A002", "2024-06-01", "Female", dx=["Membranous nephropathy"]), "u")
    db.save_record(make_record(db, "A003", "2024-07-01", "Male", birth="2010-01-01"), "u")

    assert len(db.query_records(year_from=2024, year_to=2024)) == 2
    assert len(db.query_records(genders=["Female"])) == 1
    assert len(db.query_records(diagnoses=["IgA nephropathy"])) == 2
    assert len(db.query_records(age_min=18)) == 2
    recs = db.query_records(year_from=2024, diagnoses=["IgA nephropathy"])
    assert [r["patient"]["chart_no"] for r in recs] == ["A003"]

    out = tmp_path / "out.xlsx"
    export_records(db.query_records(), str(out), {"年份": "全部"})
    wb = load_workbook(out)
    ws = wb["Biopsies"]
    headers = [c.value for c in ws[1]]
    assert ws.max_row == 4
    row = {h: c.value for h, c in zip(headers, ws[2])}
    assert row["Chart No."] == "A001"
    assert row["Age"] == 53
    assert row["Cr"] == 2.5  # 取切片當天的抽血
    assert row["Dx: IgA nephropathy"] == 1
    assert wb["Labs_all"].max_row == 7

    # 只匯出部分 section
    export_records(db.query_records(), str(out), {}, sections=["pathology"])
    headers = [c.value for c in load_workbook(out)["Biopsies"][1]]
    assert "LM" in headers and "Cr" not in headers


def test_pick_lab_row():
    rows = [{"lab_date": "2024-01-01"}, {"lab_date": "2024-02-01"}, {"lab_date": "2024-03-01"}]
    assert pick_lab_row(rows, "2024-02-01")["lab_date"] == "2024-02-01"
    assert pick_lab_row(rows, "2024-02-15") is None  # 切片當天沒抽血 → 留空
    assert pick_lab_row(rows, "") is None
    assert pick_lab_row([], "2024-01-01") is None


def test_backup(db):
    dest = db.auto_backup()
    assert dest and os.path.exists(dest)
    assert db.auto_backup() is None  # 同一天只備份一次
