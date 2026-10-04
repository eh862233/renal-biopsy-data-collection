"""讀取 TSN 腎臟疾病登錄系統匯出的 Excel，轉成可匯入資料庫的格式。

* 以標題名稱找欄位（欄位順序變動也能讀）
* 每一列 = 一次切片；檢驗數據記在切片日期（TSN 檢驗表單本身沒有抽血日期）
* 病理診斷對應到本系統的診斷清單；對應不到的放進「其他」，原始文字另存在 TSN registry 頁
"""
from datetime import date, datetime
from typing import Dict, List, Tuple

from openpyxl import load_workbook

from . import schema
from .utils import normalize_date, normalize_national_id

SHEET_NAME = "全部病患"

GENDER = {"男": "Male", "女": "Female", "M": "Male", "F": "Female",
          "Male": "Male", "Female": "Female"}

PATH_DX_MAP = {
    "Minimal change disease": "Minimal change disease",
    "Diabetic nephropathy": "Diabetic nephropathy",
    "IgA nephropathy": "IgA nephropathy",
    "Membranous glomerulonephritis": "Membranous nephropathy",
    "Membranous nephropathy": "Membranous nephropathy",
    "Focal segmental glomerulosclerosis": "FSGS",
    "Chronic tubulointerstitial nephropathy": "Chronic tubulointerstitial nephritis",
    "Acute tubulointerstitial nephropathy": "Acute interstitial nephritis",
    "Acute tubular necrosis": "Acute tubular injury / necrosis",
    "Crescentic glomerulonephritis type 1": "Crescentic GN type 1 (anti-GBM)",
    "Crescentic glomerulonephritis type 2": "Crescentic GN type 2 (immune complex)",
    "Crescentic glomerulonephritis type 3": "Crescentic GN type 3 (pauci-immune)",
    "Membranoproliferative glomerulonephritis type 1 or type 3": "MPGN",
    "Membranoproliferative glomerulonephritis type 2": "MPGN",
    "Mesangial proliferative glomerulonephritis other than IgA nephropathy":
        "Mesangial proliferative GN (non-IgA)",
    "Lupus nephritis": "Lupus nephritis (class not specified)",
    "Hypertensive nephrosclerosis": "Hypertensive nephrosclerosis",
    "Sclerosing glomerulonephritis": "Advanced chronic sclerosing nephropathy",
    "Thin basement membrane disease": "Alport syndrome / Thin basement membrane",
    "Renal transplantation": "Transplant kidney biopsy",
    "Amyloidosis": "Amyloidosis",
}

# TSN 欄位 → 切片資料欄位
DATA_COLS = {
    "系統病患ID": "tsn_patient_id",
    "登錄類別": "tsn_category",
    "醫院名稱": "tsn_hospital",
    "DM": "hx_dm",
    "HTN": "hx_htn",
    "發病原因": "tsn_etiology",
    "家族史": "tsn_family_history",
    "危險因子": "tsn_risk_factors",
    "用藥習慣": "tsn_med_habit",
    "切片編號": "tsn_biopsy_no",
    "檢驗編號": "tsn_lab_no",
    "檢驗數據上傳日期": "tsn_lab_upload_date",
    "Clinical diagnosis (Others 說明)": "clinical_dx_other",
    "Pathology diagnosis": "tsn_path_dx",
    "Pathology diagnosis (Others 說明)": "tsn_path_dx_other",
    "24hr urine protein": "urine_24hr_protein",
    "24hr urine albumin": "urine_24hr_albumin",
    "Spot UPCR": "upcr",
    "Spot UACR": "uacr",
    "U/A RBC": "ua_rbc",
    "Dysmorphic RBC": "dysmorphic_rbc",
    "ANA": "ab_ana",
    "Anti-dsDNA": "ab_dsdna",
    "ANCA": "ab_anca",
    "ANCA-MPO": "ab_mpo",
    "ANCA-PR3": "ab_pr3",
    "Anti-GBM Ab": "ab_gbm",
    "Anti-PLA2R (IgG) Ab": "ab_pla2r",
    "Cryoglobulin": "ab_cryo",
    "PEP serum": "electrophoresis",
}

# TSN 欄位 → (抽血 panel, 項目)
LAB_COLS = {
    "Creatinine": ("chem", "Cr"),
    "Blood UN": ("chem", "BUN"),
    "Cr/eGFR (系統計算)": ("chem", "eGFR"),
    "Hemoglobin": ("cbc", "Hb"),
    "Platelet": ("cbc", "PLT"),
    "Haptoglobin": ("cbc", "Haptoglobin"),
    "Serum albumin": ("liver", "Alb"),
    "AST": ("liver", "AST"),
    "ALT": ("liver", "ALT"),
    "C3": ("ig", "C3"),
    "C4": ("ig", "C4"),
    "IgG": ("ig", "IgG"),
    "IgG4": ("ig", "IgG4"),
    "IgA": ("ig", "IgA"),
    "IgM": ("ig", "IgM"),
    "IgE": ("ig", "IgE"),
    "Kappa": ("ig", "Kappa"),
    "Lambda": ("ig", "Lambda"),
    "FLC ratio": ("ig", "Kappa/Lambda"),
    "HBsAg": ("endo", "HBsAg"),
    "Anti-HCV": ("endo", "Anti-HCV"),
    "RPR/VDRL": ("endo", "VDRL"),
    "TPHA": ("endo", "TPHA"),
    "Anti-HIV": ("endo", "HIV"),
    "ASLO": ("endo", "ASLO"),
}


def _text(v) -> str:
    """儲存格值轉成文字；公式（未計算）忽略。"""
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).strip()
    return "" if s.startswith("=") else s


def _split(v: str) -> List[str]:
    return [x.strip() for x in v.split(";") if x.strip()]


def _is_number(s: str) -> bool:
    try:
        float(s.replace(",", ""))
        return True
    except ValueError:
        return False


def _find_header(ws) -> Tuple[int, Dict[str, int]]:
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        names = [str(v).strip() if v is not None else "" for v in row]
        if "身分證字號" in names:
            return r, {n: i for i, n in enumerate(names) if n}
    raise ValueError("找不到標題列（需包含「身分證字號」欄位）。這不是 TSN 匯出的 Excel 嗎？")


def parse_tsn(path: str) -> Tuple[List[Dict], List[str], List[str]]:
    """回傳 (items, skipped, warnings)。items 可直接交給 Database.import_items。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
        else:
            ws = next((w for w in wb.worksheets if any(
                "身分證字號" in [str(v) for v in row]
                for row in w.iter_rows(max_row=10, values_only=True))), None)
            if ws is None:
                raise ValueError("找不到含有「身分證字號」欄位的工作表。")
        header_row, cols = _find_header(ws)
        for need in ("身分證字號", "切片日期"):
            if need not in cols:
                raise ValueError(f"缺少必要欄位「{need}」。")

        number_fields = {f.key for f in schema.all_biopsy_fields() if f.type == schema.NUMBER}
        items, skipped, warnings = [], [], []
        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                header_row + 1):
            def get(name):
                i = cols.get(name)
                return _text(row[i]) if i is not None and i < len(row) else ""

            nid = normalize_national_id(get("身分證字號"))
            name = get("姓名")
            raw_date = get("切片日期")
            if not nid and not raw_date and not name:
                continue  # 空白列
            who = f"第 {r} 列 {name}".rstrip()
            if not nid:
                skipped.append(f"{who}：沒有身分證字號")
                continue
            bdate = normalize_date(raw_date)
            if not bdate:
                skipped.append(f"{who}：沒有切片日期或格式錯誤（{raw_date or '空白'}）")
                continue

            patient = {
                "national_id": nid,
                "chart_no": get("病歷號"),
                "name": name,
                "gender": GENDER.get(get("性別"), ""),
                "birth_date": normalize_date(get("生日")) or "",
            }

            data = {}
            for col, key in DATA_COLS.items():
                v = get(col)
                if not v:
                    continue
                if key in number_fields and not _is_number(v):
                    warnings.append(f"{who}：{col}「{v}」不是數字，未匯入")
                    continue
                if key == "tsn_lab_upload_date":
                    v = normalize_date(v) or ""
                data[key] = v
            # Clinical diagnosis（可複選）
            clin = _split(get("Clinical diagnosis"))
            known = [c for c in clin if c in schema.CLINICAL_DIAGNOSES]
            unknown = [c for c in clin if c not in schema.CLINICAL_DIAGNOSES]
            if known:
                data["clinical_dx"] = known
            if unknown:
                data["clinical_dx_other"] = "; ".join(
                    x for x in [data.get("clinical_dx_other", "")] + unknown if x)

            # 病理診斷
            diagnoses, other = [], []
            for dx in _split(data.get("tsn_path_dx", "")):
                mapped = PATH_DX_MAP.get(dx)
                if mapped:
                    if mapped not in diagnoses:
                        diagnoses.append(mapped)
                elif dx.lower() in ("others", "other"):
                    other.append(data.get("tsn_path_dx_other") or "Others")
                elif dx in schema.DIAGNOSES:
                    diagnoses.append(dx)
                else:
                    other.append(dx)

            labs: Dict[str, Dict[str, str]] = {}
            for col, (panel, analyte) in LAB_COLS.items():
                v = get(col)
                if v:
                    labs.setdefault(panel, {})[analyte] = v

            cr, bun = labs.get("chem", {}).get("Cr"), labs.get("chem", {}).get("BUN")
            if cr and bun and _is_number(cr) and _is_number(bun) and float(cr) > float(bun):
                warnings.append(f"{who}：Creatinine ({cr}) 大於 BUN ({bun})，"
                                "可能欄位對調，請匯入後核對")

            items.append({
                "row": r,
                "patient": patient,
                "biopsy_date": bdate,
                "data": data,
                "diagnoses": diagnoses,
                "diagnosis_other": "; ".join(other),
                "labs": labs,
            })
        return items, skipped, warnings
    finally:
        wb.close()
