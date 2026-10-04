"""匯出 Excel。

工作表：
  Biopsies   每次切片一列；抽血數值只取「切片日當天」的那一次（當天沒抽則留空）
  Labs_all   所有抽血紀錄（每個日期一列），方便做追蹤分析
  篩選條件   匯出時間、匯出者與篩選條件
"""
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema
from .utils import calc_age, calc_bmi, now_str, to_number

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def pick_lab_row(rows: List[Dict], biopsy_date: str) -> Optional[Dict]:
    """取切片日當天的抽血紀錄；當天沒有就回傳 None（Excel 留空）。"""
    if not biopsy_date:
        return None
    same_day = [r for r in rows if r.get("lab_date") == biopsy_date]
    return same_day[-1] if same_day else None


def _dx_col(dx: str) -> str:
    return "Dx: " + dx


def build_columns(sections: Optional[List[str]] = None):
    """回傳 [(欄位標題, 取值函式)]。sections 為要匯出的 section key，None 表示全部。"""
    cols = [
        ("Chart No.", lambda r: r["patient"]["chart_no"]),
        ("Birth date", lambda r: r["patient"].get("birth_date", "")),
        ("Gender", lambda r: r["patient"].get("gender", "")),
        ("Age", lambda r: to_number(calc_age(r["patient"].get("birth_date", ""),
                                             r["biopsy"]["data"].get("biopsy_date", "")))),
        ("Biopsy date", lambda r: r["biopsy"]["data"].get("biopsy_date", "")),
        ("Diagnosis", lambda r: "; ".join(r["biopsy"]["diagnoses"]
                                          + ([r["biopsy"]["diagnosis_other"]]
                                             if r["biopsy"].get("diagnosis_other") else []))),
    ]
    want = set(sections) if sections else None

    if want is None or "profile" in want:
        cols += [
            ("Past history", lambda r: r["patient"].get("past_history", "")),
            ("Operation history", lambda r: r["patient"].get("operation_history", "")),
        ]

    for sec in schema.SECTIONS:
        if want is not None and sec.key not in want:
            continue
        for f in sec.fields:
            if f.key in ("biopsy_date", "age"):
                continue
            if f.key == "pe_bmi":
                cols.append(("BMI", lambda r: to_number(calc_bmi(
                    r["biopsy"]["data"].get("pe_bw", ""), r["biopsy"]["data"].get("pe_height", "")))))
                continue
            title = f.label + (f" ({f.unit})" if f.unit else "")
            if f.type == schema.CHECKS:
                for opt in f.options:
                    cols.append((f"{f.label}: {opt}",
                                 lambda r, k=f.key, o=opt: 1 if o in (r["biopsy"]["data"].get(k) or []) else 0))
            elif f.type == schema.NUMBER:
                cols.append((title, lambda r, k=f.key: to_number(r["biopsy"]["data"].get(k, ""))))
            else:
                cols.append((title, lambda r, k=f.key: r["biopsy"]["data"].get(k, "")))
        for panel in sec.lab_panels:
            ptitle, analytes = schema.LAB_PANELS[panel]

            def picked(r, p=panel):
                rows = [l for l in r["labs"] if l["panel"] == p]
                return pick_lab_row(rows, r["biopsy"]["data"].get("biopsy_date", ""))

            cols.append((f"{panel.upper()} date",
                         lambda r, pk=picked: (pk(r) or {}).get("lab_date", "")))
            for a in analytes:
                cols.append((a, lambda r, pk=picked, a=a: to_number(((pk(r) or {}).get("values") or {}).get(a, ""))))
        if sec.images:
            cols.append(("Ultrasound images (n)", lambda r: len(r.get("images", []))))
        if sec.diagnosis:
            for dx in schema.DIAGNOSES:
                cols.append((_dx_col(dx), lambda r, d=dx: 1 if d in r["biopsy"]["diagnoses"] else 0))
            cols.append(("Dx: Other", lambda r: r["biopsy"].get("diagnosis_other", "")))

    cols += [
        ("Last updated", lambda r: r["biopsy"].get("updated_at", "")),
        ("Updated by", lambda r: r["biopsy"].get("updated_by", "")),
    ]
    return cols


def _write_header(ws, headers):
    ws.append(headers)
    for i, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def _autosize(ws, max_width=40):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col[:200])
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), max_width)


def export_records(records: List[Dict], path: str, criteria: Dict[str, str],
                   sections: Optional[List[str]] = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Biopsies"
    cols = build_columns(sections)
    _write_header(ws, [c[0] for c in cols])
    for rec in records:
        ws.append([fn(rec) for _, fn in cols])
    _autosize(ws)

    ws2 = wb.create_sheet("Labs_all")
    analytes = []
    for _, (_, names) in schema.LAB_PANELS.items():
        analytes += names
    _write_header(ws2, ["Chart No.", "Biopsy date", "Panel", "Lab date"] + analytes)
    for rec in records:
        for lab in sorted(rec["labs"], key=lambda l: (l["panel"], l.get("lab_date", ""))):
            names = schema.LAB_PANELS.get(lab["panel"], ("", []))[1]
            row = [rec["patient"]["chart_no"], rec["biopsy"]["data"].get("biopsy_date", ""),
                   schema.LAB_PANELS.get(lab["panel"], (lab["panel"],))[0], lab.get("lab_date", "")]
            row += [to_number(lab["values"].get(a, "")) if a in names else None for a in analytes]
            ws2.append(row)
    _autosize(ws2, 16)

    ws3 = wb.create_sheet("篩選條件")
    ws3.append(["匯出時間", now_str()])
    for k, v in criteria.items():
        ws3.append([k, v])
    ws3.append(["筆數", len(records)])
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 60

    wb.save(path)
